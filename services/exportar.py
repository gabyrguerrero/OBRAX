"""Informes diarios, semanales, mensuales y finales (CSV, Excel y PDF)."""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from .calculos import Control, periodo
from .formato import dinero, fmt

TITULOS = {"diario": "Informe diario", "semanal": "Informe semanal", "mensual": "Informe mensual", "final": "Informe final"}


@dataclass
class Tabla:
    titulo: str
    encabezados: list[str]
    filas: list[list]


@dataclass
class Informe:
    titulo: str
    subtitulo: str
    resumen: list[tuple[str, str]]
    tablas: list[Tabla]
    fotos: list[tuple[str, str]] = field(default_factory=list)  # (ruta, leyenda)


def construir(c: Control, tipo: str, ref: date) -> Informe:
    rs = periodo(c.reportes, tipo, ref)
    acum, dias = c.calcular(rs)
    recargo = 1 + c.obra.recargo_extra / 100
    actividades = [
        [a.nombre, a.unidad, fmt(acum[a.id].cantidad), fmt(c.acum[a.id].cantidad),
         fmt(c.porcentaje(a)) + "%", dinero(acum[a.id].costo)]
        for a in c.hojas()
    ]
    personal = []
    for t in c.trabajadores.values():
        ps = [p for r in rs for p in r.asistencias if p.trabajador_id == t.id]
        if ps:
            costo = sum(t.valor_hora * (p.horas + p.horas_extra * recargo) for p in ps)
            personal.append([t.nombre, t.cargo, len(ps), fmt(sum(p.horas for p in ps)),
                             fmt(sum(p.horas_extra for p in ps)), dinero(costo)])
    novedades = [[r.fecha.isoformat(), r.condiciones, r.novedades, r.observaciones]
                 for r in rs if r.novedades or r.observaciones]
    fotos = [(f.ruta, f"{c.actividad(f.actividad_id).nombre} - {r.fecha:%d/%m/%Y}")
             for r in rs for f in r.fotos if c.actividad(f.actividad_id)]
    sub = f"{rs[0].fecha:%d/%m/%Y} a {rs[-1].fecha:%d/%m/%Y}" if rs else "Sin reportes en el periodo"
    return Informe(
        f"{TITULOS[tipo]} - {c.obra.nombre}", sub,
        [("Avance acumulado de la obra", fmt(c.avance_obra()) + "%"),
         ("Costo de mano de obra del periodo", dinero(sum(d.costo for d in dias.values()))),
         ("Horas trabajadas", fmt(sum(d.horas for d in dias.values()))),
         ("Jornales (trabajador-día)", str(sum(d.personas for d in dias.values())))],
        [Tabla("Cantidades ejecutadas", ["Actividad", "Und", "Cant. periodo", "Cant. acumulada", "% avance", "Costo MO periodo"], actividades),
         Tabla("Personal y costos", ["Trabajador", "Cargo", "Días", "Horas", "Horas extra", "Costo"], personal),
         Tabla("Novedades y observaciones", ["Fecha", "Condiciones", "Novedades del personal", "Observaciones"], novedades)],
        fotos,
    )


def a_csv(inf: Informe, ruta: Path) -> None:
    """CSV con ';' y UTF-8 con BOM: Excel en español lo abre directamente."""
    with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([inf.titulo]); w.writerow([inf.subtitulo]); w.writerows(inf.resumen)
        for t in inf.tablas:
            w.writerow([]); w.writerow([t.titulo]); w.writerow(t.encabezados); w.writerows(t.filas)


def a_excel(inf: Informe, ruta: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([inf.titulo]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([inf.subtitulo]); ws.append([])
    for fila in inf.resumen:
        ws.append(list(fila))
    ws.column_dimensions["A"].width = 40
    for t in inf.tablas:
        hoja = wb.create_sheet(t.titulo[:31])
        hoja.append(t.encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for fila in t.filas:
            hoja.append(fila)
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = 22
    wb.save(ruta)


def _l1(texto) -> str:  # las fuentes base del PDF solo admiten latin-1
    return str(texto).encode("latin-1", "replace").decode("latin-1")


def a_pdf(inf: Informe, ruta: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    siguiente = {"new_x": "LMARGIN", "new_y": "NEXT"}
    pdf.set_font("Helvetica", "B", 16); pdf.cell(0, 10, _l1(inf.titulo), **siguiente)
    pdf.set_font("Helvetica", size=10); pdf.cell(0, 6, _l1(inf.subtitulo), **siguiente)
    for etiqueta, valor in inf.resumen:
        pdf.cell(0, 6, _l1(f"{etiqueta}: {valor}"), **siguiente)
    for t in inf.tablas:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 12); pdf.cell(0, 8, _l1(t.titulo), **siguiente)
        pdf.set_font("Helvetica", size=8)
        if not t.filas:
            pdf.cell(0, 6, "Sin datos en este periodo", **siguiente)
            continue
        with pdf.table() as tabla:
            fila = tabla.row()
            for h in t.encabezados:
                fila.cell(_l1(h))
            for datos in t.filas:
                fila = tabla.row()
                for x in datos:
                    fila.cell(_l1(x))
    if inf.fotos:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 12); pdf.cell(0, 8, "Fotografias", **siguiente)
        pdf.set_font("Helvetica", size=8)
        for foto, leyenda in inf.fotos:
            try:
                pdf.image(foto, w=70)
                pdf.cell(0, 5, _l1(leyenda), **siguiente)
            except Exception:
                pdf.cell(0, 5, _l1(f"(no se pudo insertar la foto: {leyenda})"), **siguiente)
    pdf.output(str(ruta))
